"""Adapter contracts: validation, full FSM semantics, and native tool round trips."""

from __future__ import annotations

import copy
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any
import unittest
from unittest.mock import Mock

HAS_PYDANTIC = importlib.util.find_spec("pydantic") is not None


def calculation_arguments(**parameters: float) -> dict[str, Any]:
    return {
        "scenario_id": "molten_salt_steam",
        "inputs": [{"name": name, "value": value} for name, value in parameters.items()],
        "targets": None,
        "solution_mode": "AUTO",
        "generate_excel": False,
    }


@unittest.skipUnless(HAS_PYDANTIC, "安装 integrations/requirements.txt 后运行接入层测试")
class TestIndustrialTools(unittest.TestCase):
    def setUp(self) -> None:
        from integrations.service import IndustrialToolService
        self.service = IndustrialToolService()

    def test_discovery_exposes_units_defaults_bounds_and_runtime_aliases(self) -> None:
        result = self.service.invoke("industrial_list_scenarios", {})
        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(
            {item["scenario_id"] for item in result["scenarios"]},
            {"molten_salt_steam", "financial_profit_model"},
        )
        for item in result["scenarios"]:
            self.assertTrue(item["description"])
            spec = self.service.invoke("industrial_get_scenario_spec", {"scenario_id": item["scenario_id"]})
            self.assertEqual(spec["status"], "SUCCESS")
            self.assertIn("parameters", spec["spec"])
            self.assertIn("alias_map", spec["spec"])
        spec = self.service.get_scenario_spec("molten_salt_steam")["spec"]
        self.assertEqual(spec["parameters"]["steam_pressure_mpa"]["canonical_unit"], "MPa")
        self.assertEqual(spec["parameters"]["heater_efficiency"]["default"], 0.985)
        self.assertEqual(spec["outputs"]["dynamic_investment_wanke"]["canonical_unit"], "万元")
        spec["parameters"]["heater_efficiency"]["default"] = 0
        self.assertEqual(self.service.get_scenario_spec("molten_salt_steam")["spec"]
                         ["parameters"]["heater_efficiency"]["default"], 0.985)

    def test_forward_preserves_all_states_audit_warnings_and_input(self) -> None:
        arguments = calculation_arguments(**{"蒸汽压力": 1.5, "蒸汽温度": 200, "蒸汽流量": 100})
        original = copy.deepcopy(arguments)
        result = self.service.invoke("industrial_calculate", arguments)
        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(result["results"]["dynamic_investment_wanke"], 22650)
        self.assertEqual([item["state"].split("_")[0] for item in result["fsm_trace"]],
                         [f"S{index}" for index in range(9)])
        self.assertTrue(result["audit"]["passed"])
        self.assertEqual(result["warnings"], result["audit"]["warnings"])
        self.assertTrue(result["warnings"])
        self.assertEqual(result["artifacts"], {})
        self.assertEqual(arguments, original)

    def test_inverse_uses_budget_and_returns_feasible_capacity(self) -> None:
        arguments = calculation_arguments(steam_pressure_mpa=1.5, steam_temperature_c=200,
                                          target_dynamic_investment_wanke=22650)
        arguments["targets"] = ["steam_flow_th"]
        result = self.service.invoke("industrial_calculate", arguments)
        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(result["solution_mode"], "INVERSE")
        self.assertGreater(result["results"]["steam_flow_th"], 100)
        self.assertLessEqual(result["results"]["dynamic_investment_wanke"], 22650 + 1e-7)
        self.assertEqual(len(result["fsm_trace"]), 9)

    def test_financial_forward_and_inverse(self) -> None:
        arguments = calculation_arguments(unit_price_cny=100, sales_volume=100000)
        arguments["scenario_id"] = "financial_profit_model"
        result = self.service.invoke("industrial_calculate", arguments)
        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(result["results"]["net_profit_wanke"], 172.5)
        arguments["inputs"] = [{"name": "unit_price_cny", "value": 100},
                               {"name": "target_net_profit_wanke", "value": 172.5}]
        arguments["targets"] = ["sales_volume"]
        result = self.service.invoke("industrial_calculate", arguments)
        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(result["results"]["sales_volume"], 100000)

    def test_missing_parameters_interrupt_without_delivery(self) -> None:
        result = self.service.invoke("industrial_calculate", calculation_arguments(steam_pressure_mpa=1.5))
        self.assertEqual(result["status"], "INTERRUPTED")
        self.assertIn("steam_flow_th", result["missing_parameters"])
        self.assertNotIn("artifacts", result)
        self.assertEqual(result["fsm_trace"][-1]["state"], "S1_CANONICAL_ALIGNMENT")

    def test_domain_failure_keeps_engine_diagnostics(self) -> None:
        result = self.service.invoke("industrial_calculate", calculation_arguments(
            steam_pressure_mpa=1.5, steam_temperature_c=100, steam_flow_th=100))
        self.assertEqual(result["status"], "FAILED")
        self.assertTrue(result["diagnostics"])
        self.assertNotIn("artifacts", result)

    def test_invalid_arguments_never_start_fsm_and_are_strict_json(self) -> None:
        from integrations.service import encode_json
        valid = calculation_arguments(steam_pressure_mpa=1.5)
        cases: list[Any] = ["{", "[]", "null", '{"scenario_id":"a","scenario_id":"b"}',
                            '{"inputs":[{"name":"pressure","value":NaN}]}']
        for value in (True, "1.5 MPa", float("nan"), float("inf"), None):
            bad = copy.deepcopy(valid)
            bad["inputs"][0]["value"] = value
            cases.append(bad)
        for extra in ({"audit_level": "OFF"}, {"options": {"generate_excel": True}},
                      {"solution_mode": "arbitrary"}, {"generate_excel": "true"}):
            cases.append({**valid, **extra})
        cases.append({**valid, "inputs": valid["inputs"] * 2})
        cases.append({**valid, "inputs": [{"name": "蒸汽压力", "value": 1.5},
                                         {"name": "steam_pressure_mpa", "value": 2}]})
        for arguments in cases:
            with self.subTest(arguments=str(arguments)):
                result = self.service.invoke("industrial_calculate", arguments)
                self.assertEqual(result["status"], "FAILED")
                self.assertEqual(result["error"]["code"], "INVALID_ARGUMENTS")
                self.assertNotIn("fsm_trace", result)
                json.loads(encode_json(result))

    def test_unknown_tool_and_unknown_scenario(self) -> None:
        self.assertEqual(self.service.invoke("__import__", {})["error"]["code"], "UNKNOWN_TOOL")
        result = self.service.invoke("industrial_get_scenario_spec", {"scenario_id": "../secret"})
        self.assertEqual(result["error"]["code"], "UNKNOWN_SCENARIO")
        result = self.service.invoke("industrial_calculate", {"scenario_id": "absent", "inputs": []})
        self.assertEqual(result["status"], "FAILED")

    def test_internal_exception_is_contained(self) -> None:
        from integrations.service import IndustrialToolService
        skill = Mock()
        skill.list_scenarios.side_effect = RuntimeError("internal sensitive detail")
        with self.assertLogs("integrations.service", level="ERROR"):
            result = IndustrialToolService(skill).invoke("industrial_list_scenarios", {})
        self.assertEqual(result["error"]["code"], "INTERNAL_ERROR")
        self.assertNotIn("sensitive", json.dumps(result))

    def test_excel_is_explicit_and_has_formula_and_cached_result(self) -> None:
        import openpyxl
        arguments = calculation_arguments(steam_pressure_mpa=1.5, steam_temperature_c=200, steam_flow_th=100)
        arguments["generate_excel"] = True
        result = self.service.invoke("industrial_calculate", arguments)
        self.assertEqual(result["status"], "SUCCESS")
        path = Path(result["artifacts"]["excel_path"])
        self.addCleanup(path.unlink, missing_ok=True)
        self.assertTrue(path.is_absolute())
        formula_book = openpyxl.load_workbook(path, data_only=False)
        values_book = openpyxl.load_workbook(path, data_only=True)
        try:
            self.assertEqual(len(formula_book.sheetnames), 7)
            self.assertTrue(formula_book["计算"]["D5"].value.startswith("="))
            self.assertIsInstance(values_book["计算"]["D5"].value, (int, float))
        finally:
            formula_book.close()
            values_book.close()


@unittest.skipUnless(HAS_PYDANTIC, "安装 integrations/requirements.txt 后运行接入层测试")
class TestNativeFunctionCalling(unittest.TestCase):
    def test_strict_schemas_for_both_openai_apis(self) -> None:
        from integrations.function_calling import get_openai_tools
        def check(node: Any) -> None:
            if isinstance(node, dict):
                if node.get("type") == "object":
                    self.assertIs(node["additionalProperties"], False)
                    self.assertEqual(set(node["required"]), set(node.get("properties", {})))
                self.assertNotIn("default", node)
                for child in node.values():
                    check(child)
            elif isinstance(node, list):
                for child in node:
                    check(child)
        for api in ("responses", "chat_completions"):
            tools = get_openai_tools(api)
            self.assertEqual(len(tools), 3)
            for tool in tools:
                function = tool if api == "responses" else tool["function"]
                self.assertTrue(function["strict"])
                check(function["parameters"])
        with self.assertRaises(ValueError):
            get_openai_tools("unknown")

    @unittest.skipUnless(importlib.util.find_spec("jsonschema"), "需要 jsonschema")
    def test_schema_accepts_numeric_inputs_and_rejects_unknown_keys(self) -> None:
        import jsonschema
        from integrations.function_calling import get_function_definitions
        schema = get_function_definitions()[2]["parameters"]
        jsonschema.Draft202012Validator.check_schema(schema)
        arguments = calculation_arguments(steam_pressure_mpa=1.5, steam_temperature_c=200, steam_flow_th=100)
        jsonschema.validate(arguments, schema)
        for bad in ({**arguments, "other": 1}, {**arguments, "targets": "steam_flow_th"},
                    {key: value for key, value in arguments.items() if key != "generate_excel"}):
            with self.assertRaises(jsonschema.ValidationError):
                jsonschema.validate(bad, schema)

    def test_responses_replays_reasoning_and_correlates_multiple_calls(self) -> None:
        from integrations.function_calling import run_responses_agent
        reasoning = SimpleNamespace(type="reasoning", id="reasoning_1", summary=[])
        first = SimpleNamespace(output=[reasoning,
            SimpleNamespace(type="function_call", name="industrial_list_scenarios", arguments="{}", call_id="first"),
            SimpleNamespace(type="function_call", name="industrial_get_scenario_spec",
                            arguments='{"scenario_id":"molten_salt_steam"}', call_id="second")])
        final = SimpleNamespace(output=[], output_text="需要蒸汽压力、温度和流量。")
        client = Mock()
        client.responses.create.side_effect = [first, final]
        result = run_responses_agent(client, "test-model", "列出计算场景")
        self.assertIs(result, final)
        request = client.responses.create.call_args_list[1].kwargs
        self.assertIs(request["input"][1], reasoning)
        outputs = request["input"][-2:]
        self.assertEqual([item["call_id"] for item in outputs], ["first", "second"])
        self.assertTrue(all(json.loads(item["output"])["status"] == "SUCCESS" for item in outputs))

    def test_responses_round_limit(self) -> None:
        from integrations.function_calling import run_responses_agent
        client = Mock()
        client.responses.create.return_value = {"output": [
            {"type": "function_call", "name": "industrial_list_scenarios", "arguments": "{}", "call_id": "call"}]}
        with self.assertRaises(RuntimeError):
            run_responses_agent(client, "test-model", "loop", max_rounds=2)
        self.assertEqual(client.responses.create.call_count, 2)

    def test_chat_completions_tool_call_id_and_failed_arguments(self) -> None:
        from integrations.function_calling import execute_chat_tool_calls
        outputs = execute_chat_tool_calls({"tool_calls": [
            {"id": "call_good", "type": "function", "function": {"name": "industrial_list_scenarios", "arguments": "{}"}},
            {"id": "call_bad", "type": "function", "function": {"name": "industrial_calculate", "arguments": "{"}},
        ]})
        self.assertEqual([item["tool_call_id"] for item in outputs], ["call_good", "call_bad"])
        self.assertEqual([json.loads(item["content"])["status"] for item in outputs], ["SUCCESS", "FAILED"])

    @unittest.skipUnless(importlib.util.find_spec("openai"), "需要 openai SDK")
    def test_real_openai_sdk_with_offline_mock_transport(self) -> None:
        import httpx
        from openai import OpenAI
        from integrations.function_calling import run_responses_agent
        requests: list[dict[str, Any]] = []

        def respond(request: httpx.Request) -> httpx.Response:
            requests.append(json.loads(request.content))
            output = ([{"type": "function_call", "id": "fc_1", "call_id": "call_1",
                        "name": "industrial_list_scenarios", "arguments": "{}", "status": "completed"}]
                      if len(requests) == 1 else [])
            return httpx.Response(200, json={"id": f"resp_{len(requests)}", "object": "response",
                "created_at": 0, "model": "offline-test-model", "status": "completed", "output": output})

        with OpenAI(api_key="offline-test-key", base_url="https://offline.invalid/v1",
                    http_client=httpx.Client(transport=httpx.MockTransport(respond))) as client:
            result = run_responses_agent(client, "offline-test-model", "列出场景")
        self.assertEqual(result.status, "completed")
        self.assertEqual(len(requests), 2)
        self.assertEqual(requests[1]["input"][-1]["call_id"], "call_1")
        self.assertEqual(json.loads(requests[1]["input"][-1]["output"])["status"], "SUCCESS")


@unittest.skipUnless(HAS_PYDANTIC and importlib.util.find_spec("langchain_core"), "需要 LangChain 可选依赖")
class TestLangChainTools(unittest.TestCase):
    def test_real_structured_tool_invocation_and_missing_inputs(self) -> None:
        from integrations.langchain_tools import get_langchain_tools
        tools = {tool.name: tool for tool in get_langchain_tools()}
        result = tools["industrial_calculate"].invoke(calculation_arguments(
            steam_pressure_mpa=1.5, steam_temperature_c=200, steam_flow_th=100))
        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(result["results"]["dynamic_investment_wanke"], 22650)
        result = tools["industrial_calculate"].invoke(calculation_arguments(steam_pressure_mpa=1.5))
        self.assertEqual(result["status"], "INTERRUPTED")


if __name__ == "__main__":
    unittest.main()
