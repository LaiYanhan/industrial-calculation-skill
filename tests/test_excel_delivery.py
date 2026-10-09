"""openpyxl 回填验收：真实七表、数组公式、格式、缓存与原子失败。"""
from contextlib import redirect_stdout
import hashlib
import io
import json
import math
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.formula import ArrayFormula

from engine.exporter.excel_exporter import ExcelTemplateExporter
from skill_api import CalculationSkill
from subagent_workspace.ci_runner import run_benchmark_suite

ROOT = Path(__file__).resolve().parents[1]
BASE = {"steam_pressure_mpa": 1.5, "steam_temperature_c": 200, "steam_flow_th": 100}


def formula_signature(value: object) -> object:
    return (value.text, value.ref) if isinstance(value, ArrayFormula) else value


class TestExcelDelivery(unittest.TestCase):
    def test_all_seven_sheets_formula_style_and_snapshot_preservation(self) -> None:
        template = ROOT / "输入输出.xlsx"
        original_hash = hashlib.sha256(template.read_bytes()).hexdigest()
        response = CalculationSkill().calculate("molten_salt_steam", BASE, options={"generate_excel": True})
        self.assertEqual(response["status"], "SUCCESS", response)
        output = Path(response["artifacts"]["excel_path"])
        self.addCleanup(output.unlink, missing_ok=True)
        self.assertEqual(output.parent, ROOT / "exports")
        self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest(), original_hash)
        source = openpyxl.load_workbook(template, data_only=False)
        written = openpyxl.load_workbook(output, data_only=False)
        values = openpyxl.load_workbook(output, data_only=True)
        self.addCleanup(source.close)
        self.addCleanup(written.close)
        self.addCleanup(values.close)
        self.assertEqual(source.sheetnames, written.sheetnames)
        self.assertEqual(len(written.sheetnames), 7)
        self.assertEqual(len(source._external_links), len(written._external_links))
        for sheet in source:
            other = written[sheet.title]
            self.assertEqual(str(sheet.merged_cells), str(other.merged_cells))
            self.assertEqual(sheet.freeze_panes, other.freeze_panes)
            self.assertEqual(sheet.sheet_state, other.sheet_state)
            for key in sheet.column_dimensions:
                self.assertEqual(sheet.column_dimensions[key].width, other.column_dimensions[key].width)
            for row in sheet:
                for cell in row:
                    copied = other[cell.coordinate]
                    self.assertEqual(cell._style, copied._style, f"{sheet.title}!{cell.coordinate}")
                    if cell.data_type == "f":
                        self.assertEqual(formula_signature(cell.value), formula_signature(copied.value))
                        cached = values[sheet.title][cell.coordinate].value
                        self.assertIsInstance(cached, (int, float), f"缺失缓存: {sheet.title}!{cell.coordinate}")
                        self.assertTrue(math.isfinite(cached))
        self.assertAlmostEqual(values["工程概算"]["G22"].value,
                               response["results"]["dynamic_investment_wanke"], delta=1e-7)
        self.assertEqual(values["选型"]["D9"].value, 8650)
        self.assertEqual(values["造价"]["E17"].value, 21640)
        self.assertAlmostEqual(values["工程概算"]["G19"].value, values["造价"]["E9"].value, delta=1e-7)
        self.assertEqual(values["计算"]["D3"].value, response["results"]["steam_enthalpy_kj_kg"])

    def test_changed_inputs_solar_and_interest_refresh_dependent_caches(self) -> None:
        response = CalculationSkill().calculate("molten_salt_steam", BASE | {
            "salt_type_id": 1, "steam_temperature_c": 400, "construction_interest_wanke": 500,
            "steam_supply_hours_h": 10, "salt_mass_margin": 1.2,
        }, options={"generate_excel": True})
        self.assertEqual(response["status"], "SUCCESS", response)
        output = Path(response["artifacts"]["excel_path"])
        self.addCleanup(output.unlink, missing_ok=True)
        formulas = openpyxl.load_workbook(output, data_only=False)
        values = openpyxl.load_workbook(output, data_only=True)
        self.addCleanup(formulas.close)
        self.addCleanup(values.close)
        self.assertEqual(values["输入"]["D7"].value, 10)
        self.assertEqual(values["选型"]["D5"].value, "太阳盐")
        self.assertEqual(formulas["选型"]["D8"].value, "=AVERAGE(熔盐参数!K11:K12)")
        self.assertAlmostEqual(values["选型"]["D8"].value, 1.51653)
        self.assertEqual(values["工程概算"]["F21"].value, 500)
        self.assertAlmostEqual(values["工程概算"]["G22"].value, response["results"]["dynamic_investment_wanke"], delta=1e-7)
        self.assertEqual(values["计算"]["D10"].value, 26.4)
        self.assertEqual(values["计算"]["J11"].value, 33)
        self.assertEqual(values["选型"]["D9"].value, response["results"]["molten_salt_mass_t"])

    def test_large_transformer_array_override(self) -> None:
        response = CalculationSkill().calculate("molten_salt_steam", BASE | {"steam_flow_th": 300},
                                                options={"generate_excel": True})
        self.assertEqual(response["status"], "SUCCESS", response)
        output = Path(response["artifacts"]["excel_path"])
        self.addCleanup(output.unlink, missing_ok=True)
        workbook = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["选型"]["D11"].value, "=SUM(120.0,120.0)")


class TestGenericExcelExporter(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(dir=ROOT / "subagent_workspace")
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.template = self.folder / "template.xlsx"
        self.output = self.folder / "output.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Test's sheet"
        sheet["A1"] = 2
        sheet["A1"].font = Font(bold=True, color="FF112233")
        sheet["A1"].fill = PatternFill("solid", fgColor="FFCCEEDD")
        sheet["B1"] = "=A1*3"
        sheet["C1"] = ArrayFormula(ref="C1", text="=SUM(A1:B1)")
        sheet.merge_cells("A3:B3")
        sheet["A3"] = "Merged"
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 21
        workbook.save(self.template)
        workbook.close()
        self.exporter = ExcelTemplateExporter(self.template)

    def test_input_write_with_explicit_formula_cache(self) -> None:
        self.exporter.export(self.output, {"'Test''s sheet'!$A$1": 4},
                             formula_values={"B1": 12, "C1": 16})
        values = openpyxl.load_workbook(self.output, data_only=True)
        formulas = openpyxl.load_workbook(self.output, data_only=False)
        self.addCleanup(values.close)
        self.addCleanup(formulas.close)
        self.assertEqual(values.active["B1"].value, 12)
        self.assertEqual(values.active["C1"].value, 16)
        self.assertEqual(formulas.active["B1"].value, "=A1*3")
        self.assertIsInstance(formulas.active["C1"].value, ArrayFormula)
        self.assertTrue(formulas.active["A1"].font.bold)
        self.assertEqual(formulas.active.freeze_panes, "B2")

    def test_failures_leave_existing_destination_unchanged(self) -> None:
        self.output.write_bytes(b"existing destination")
        for mapping in ({"Missing!A1": 4}, {"B1": 5}, {"C1": 5}, {"B3": 1},
                        {"XFE1": 1}, {"A0": 1}, {"A1:A2": 1}, {"A1": math.nan}):
            with self.subTest(mapping=mapping), self.assertRaises(ValueError):
                self.exporter.export(self.output, mapping)
            self.assertEqual(self.output.read_bytes(), b"existing destination")
        with self.assertRaises(ValueError):
            self.exporter.export(self.template, {"A1": 4})
        with self.assertRaises(ValueError):
            self.exporter.export(self.output, {}, formula_values={"A1": 2})
        with self.assertRaises(ValueError):
            self.exporter.export(self.output, {}, formula_values={"B1": math.nan})
        self.assertEqual(self.output.read_bytes(), b"existing destination")
        self.assertEqual(sorted(path.name for path in self.folder.iterdir()), ["output.xlsx", "template.xlsx"])

    def test_saving_error_does_not_create_a_placeholder(self) -> None:
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.exporter.export(self.output, {"A1": 4})
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.folder.iterdir()), [self.template])

    def test_explicit_override_and_bare_output_filename(self) -> None:
        self.exporter.export(self.output, {"B1": "=A1*4"}, formula_overrides={"B1"}, formula_values={"B1": 8})
        workbook = openpyxl.load_workbook(self.output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["B1"].value, "=A1*4")


class TestBenchmarkGate(unittest.TestCase):
    def test_inverse_placeholder_cannot_pass(self) -> None:
        fake_case = [{"case_id": "broken_inverse", "mode": "INVERSE", "target_param": "steam_flow_th",
                      "given": {"target_dynamic_investment_wanke": 22650}, "expected_range": [95,105]}]
        with patch("subagent_workspace.ci_runner.json.load", return_value=fake_case), \
             patch("subagent_workspace.ci_runner.CalculationSkill.calculate", return_value={"status":"FAILED"}), \
             redirect_stdout(io.StringIO()):
            self.assertFalse(run_benchmark_suite("molten_salt_steam"))

    def test_nonfinite_forward_result_cannot_pass(self) -> None:
        fake_case = [{"mode": "FORWARD", "inputs": {}, "expected": {"value": 100}}]
        with patch("subagent_workspace.ci_runner.json.load", return_value=fake_case), \
             patch("subagent_workspace.ci_runner.CalculationSkill.calculate",
                   return_value={"status":"SUCCESS", "results":{"value":float("nan")}}), \
             redirect_stdout(io.StringIO()):
            self.assertFalse(run_benchmark_suite("molten_salt_steam"))
